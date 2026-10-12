#!/usr/bin/env python3
"""Generate the Growth Play slide from a JSON spec.

Writes a Google Slides-ready PowerPoint file and a PNG preview.
Layout matches the 1920×1080 HTML artboard kept as the template.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
FONTS = ROOT / "fonts"
TERRA_ZIP = Path.home() / "Downloads" / "Terra — Brand Design System.zip"

W, H = 1920, 1080
EMU_PER_PX = 6350
PAD_X = 72
HEADER_BOTTOM = 64 + 108
MAIN_Y = HEADER_BOTTOM + 40
COL_W = (W - PAD_X * 2 - 28) / 2
ROW_BUDGET = 748 - 28
ROW1 = ROW_BUDGET * 1.3 / 2.3
ROW2 = ROW_BUDGET * 1.0 / 2.3
LEFT = PAD_X
RIGHT = PAD_X + COL_W + 28
ROW2_Y = MAIN_Y + ROW1 + 28

OAT = (0xFE, 0xFD, 0xFB)
SLAB = (0xF4, 0xF3, 0xF0)
INK = (0x1B, 0x1A, 0x18)
BLACK = (0x00, 0x00, 0x00)
MUTED = (0x7B, 0x79, 0x74)
BLUE_SLAB = (0xBE, 0xDF, 0xFE)
NAVY = (0x00, 0x14, 0x33)
BULLET = (0x39, 0x5A, 0xFA)
PLACE = (0x4A, 0x56, 0x70)

PANELS = (
    {"key": "who", "heading": "Who we target", "icon": "Signal-Intent.png", "fill": SLAB, "ink": INK, "x": LEFT, "y": MAIN_Y, "h": ROW1, "gap": 12},
    {"key": "considerations", "heading": "Considerations", "icon": "icon-puzzle.png", "fill": SLAB, "ink": NAVY, "x": LEFT, "y": ROW2_Y, "h": ROW2, "gap": 12},
    {"key": "talk", "heading": "Talk track", "icon": "Outbound.png", "fill": BLUE_SLAB, "ink": NAVY, "x": RIGHT, "y": MAIN_Y, "h": ROW1 + 28 + ROW2, "gap": 18},
)


def emu(px):
    return Emu(int(round(px * EMU_PER_PX)))


def rgb(color):
    return RGBColor(*color)


def is_prompt(text):
    text = text.strip()
    return len(text) >= 2 and text.startswith("[") and text.endswith("]")


def load_spec(path):
    path = Path(path)
    if path.is_dir():
        path = path / "overview.md"
    if path.suffix.lower() == ".md":
        return load_markdown(path)
    data = json.loads(path.read_text())
    title = str(data.get("title", "")).strip()
    if not title:
        raise SystemExit("The slide file needs a title.")
    go_live = str(data.get("go_live") or "[Date]").strip() or "[Date]"
    who = lines_of(data.get("who"), "who")
    talk = [str(item).strip() for item in data.get("talk") or []]
    talk = [item for item in talk if item]
    considerations = lines_of(data.get("considerations"), "considerations")
    if not who or not talk or not considerations:
        raise SystemExit("who, talk, and considerations each need at least one line.")
    return {
        "title": title,
        "go_live": go_live,
        "who": who,
        "why": why,
        "talk": talk,
        "considerations": considerations,
    }


def lines_of(rows, name):
    if not isinstance(rows, list):
        raise SystemExit(f"{name} must be a list.")
    parsed = []
    for row in rows:
        if isinstance(row, str) and row.strip():
            parsed.append(row.strip())
            continue
        if isinstance(row, dict) and str(row.get("label", "")).strip() and str(row.get("text", "")).strip():
            parsed.append((str(row["label"]).strip(), str(row["text"]).strip()))
            continue
        raise SystemExit(f"Each {name} item needs text, or a label and text.")
    return parsed


SLIDE_KEYS = {
    "who we target": "who",
    "talk track": "talk",
    "considerations": "considerations",
}


def load_markdown(path):
    text = Path(path).read_text()
    match = re.search(r"^## Slide\s*$([\s\S]*?)(?=^## |\Z)", text, re.M)
    if not match:
        raise SystemExit(f"{path} has no ## Slide section.")
    body = match.group(1)
    title = markdown_field(body, "Title")
    if not title:
        raise SystemExit(f"{path} ## Slide needs a Title line.")
    go_live = markdown_field(body, "Go-live") or "[Date]"
    blocks = {key: [] for key in ("who", "talk", "considerations")}
    current = None
    for line in body.splitlines():
        heading = re.match(r"^###\s+(.+?)\s*$", line)
        if heading:
            current = SLIDE_KEYS.get(heading.group(1).strip().lower())
            if current is None:
                raise SystemExit(f"Unknown slide block: {heading.group(1)}")
            continue
        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if not bullet or current is None:
            continue
        item = bullet.group(1).strip()
        labeled = re.match(r"^\*\*(.+?):\*\*\s*(.+)$", item)
        if labeled and current in ("who", "considerations"):
            blocks[current].append((labeled.group(1).strip(), labeled.group(2).strip()))
        elif current in ("talk", "considerations", "who"):
            blocks[current].append(item)
        else:
            raise SystemExit(f"Unknown slide bullet in {current}: {path}")
    spec = {"title": title, "go_live": go_live, **blocks}
    missing = [name for name in ("who", "talk", "considerations") if not spec[name]]
    if missing:
        raise SystemExit(f"## Slide is missing {', '.join(missing)}.")
    return spec


def markdown_field(body, name):
    match = re.search(rf"^{re.escape(name)}:\s*(.+?)\s*$", body, re.M)
    return match.group(1).strip() if match else ""


def ensure_roobert():
    regular = FONTS / "RoobertVF.ttf"
    mono = FONTS / "RoobertSemiMonoVF.ttf"
    if regular.exists() and mono.exists():
        return regular, mono
    if not TERRA_ZIP.exists():
        return None
    cache = Path("/tmp/growth-play-slide-fonts")
    cache.mkdir(parents=True, exist_ok=True)
    regular = cache / "RoobertVF.ttf"
    mono = cache / "RoobertSemiMonoVF.ttf"
    if not regular.exists() or not mono.exists():
        with zipfile.ZipFile(TERRA_ZIP) as archive:
            regular.write_bytes(archive.read("fonts/RoobertVF.ttf"))
            mono.write_bytes(archive.read("fonts/RoobertSemiMonoVF.ttf"))
    return regular, mono


def font(path, px, weight):
    face = ImageFont.truetype(path, px)
    if hasattr(face, "set_variation_by_axes"):
        try:
            face.set_variation_by_axes([weight, 0])
        except OSError:
            pass
    return face


def poppins(weight):
    name = {400: "Poppins-Regular.ttf", 500: "Poppins-Medium.ttf", 600: "Poppins-SemiBold.ttf", 700: "Poppins-Bold.ttf"}[weight]
    return FONTS / name


def tracked_width(text, face, tracking_em):
    extra = tracking_em * face.size
    width = sum(face.getlength(ch) for ch in text)
    if len(text) > 1:
        width += extra * (len(text) - 1)
    return width


def draw_tracked(draw, xy, text, face, fill, tracking_em):
    x, y = xy
    extra = tracking_em * face.size
    for index, ch in enumerate(text):
        draw.text((x, y), ch, font=face, fill=fill)
        x += face.getlength(ch)
        if index < len(text) - 1:
            x += extra
    return x


def wrap(text, face, max_width):
    lines, current = [], ""
    for word in text.split():
        trial = word if not current else f"{current} {word}"
        if face.getlength(trial) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def content_of(spec, panel):
    return spec[panel["key"]]


def labeled_lines(label, text, body, strong, max_width):
    prefix = f"{label}: "
    first_width = max_width - strong.getlength(prefix)
    if first_width <= 0:
        return [("", text)]
    words = text.split()
    first, rest = "", []
    for index, word in enumerate(words):
        trial = word if not first else f"{first} {word}"
        if body.getlength(trial) <= first_width:
            first = trial
        else:
            rest = words[index:]
            break
    lines = [(prefix, first)]
    if rest:
        lines.extend(("", line) for line in wrap(" ".join(rest), body, max_width))
    return lines


def item_lines(item, body, strong, max_width):
    if isinstance(item, tuple):
        label, text = item
        lines = labeled_lines(label, text, body, strong, max_width)
        if len(lines) > 4:
            raise SystemExit(f"{label} needs more than four lines. Shorten it before generating.")
        return lines
    return wrap(item, body, max_width)


def assert_fits(spec, body, strong):
    max_width = COL_W - 40 - 10 - 18 - 40
    for panel in PANELS:
        cursor = panel["y"] + 32 + 80 + 22
        for item in content_of(spec, panel):
            lines = len(item_lines(item, body, strong, max_width))
            cursor += 36 * lines + panel["gap"]
        if cursor - panel["gap"] > panel["y"] + panel["h"] - 12:
            raise SystemExit(
                f"{panel['heading']} does not fit on the template. Shorten it before generating."
            )


def build_png(spec, out_path, sans, mono):
    image = Image.new("RGB", (W, H), OAT)
    draw = ImageDraw.Draw(image)
    eyebrow = font(mono, 20, 600)
    title = font(sans, 80, 600)
    live_label = font(mono, 18, 600)
    live_date = font(sans, 26, 500)
    heading = font(sans, 40, 600)
    body = font(sans, 28, 400)
    strong = font(sans, 28, 600)
    assert_fits(spec, body, strong)

    stack_h = 20 + 14 + 76
    eyebrow_y = HEADER_BOTTOM - stack_h
    draw_tracked(draw, (PAD_X, eyebrow_y), "GROWTH PLAY", eyebrow, MUTED, 0.12)
    draw_tracked(draw, (PAD_X, eyebrow_y + 34), spec["title"], title, BLACK, -0.04)

    label_w = tracked_width("GO-LIVE", live_label, 0.12)
    date_w = live_date.getlength(spec["go_live"])
    pill_w = 26 + label_w + 14 + date_w + 26
    pill_h = 14 + 31 + 14
    pill_x = W - PAD_X - pill_w
    pill_y = HEADER_BOTTOM - 4 - pill_h
    draw.rounded_rectangle([pill_x, pill_y, pill_x + pill_w, pill_y + pill_h], radius=pill_h / 2, fill=SLAB)
    cursor = draw_tracked(draw, (pill_x + 26, pill_y + 16), "GO-LIVE", live_label, MUTED, 0.12)
    draw.text((cursor + 14, pill_y + 14), spec["go_live"], font=live_date, fill=INK)

    max_width = COL_W - 40 - 10 - 18 - 40
    for panel in PANELS:
        x, y, h = panel["x"], panel["y"], panel["h"]
        draw.rounded_rectangle([x, y, x + COL_W, y + h], radius=24, fill=panel["fill"])
        glyph = Image.open(ASSETS / "icons" / panel["icon"]).convert("RGBA").resize((80, 80), Image.Resampling.LANCZOS)
        image.paste(glyph, (int(x + 40), int(y + 32)), glyph)
        draw_tracked(draw, (x + 140, y + 50), panel["heading"], heading, panel["ink"], -0.025)
        text_x = x + 68
        cursor_y = y + 134
        for item in content_of(spec, panel):
            draw.rounded_rectangle([x + 40, cursor_y + 13, x + 50, cursor_y + 23], radius=3, fill=BULLET)
            lines = item_lines(item, body, strong, max_width)
            if isinstance(item, tuple):
                label, text = item
                rest_color = PLACE if is_prompt(text) else panel["ink"]
                for index, (prefix, line) in enumerate(lines):
                    line_y = cursor_y + index * 36
                    if prefix:
                        draw.text((text_x, line_y), prefix, font=strong, fill=panel["ink"])
                        draw.text((text_x + strong.getlength(prefix), line_y), line, font=body, fill=rest_color)
                    else:
                        draw.text((text_x, line_y), line, font=body, fill=rest_color)
            else:
                color = PLACE if is_prompt(item) else panel["ink"]
                for index, line in enumerate(lines):
                    draw.text((text_x, cursor_y + index * 36), line, font=body, fill=color)
            cursor_y += 36 * len(lines) + panel["gap"]

    logo = Image.open(ASSETS / "logo" / "Clay_Logo_3D_Blk.png").convert("RGBA")
    logo_h = 40
    logo_w = round(logo.width * logo_h / logo.height)
    logo = logo.resize((logo_w, logo_h), Image.Resampling.LANCZOS)
    image.paste(logo, (PAD_X, 1000 + (120 - 40) // 2), logo)
    image.save(out_path, "PNG")


def set_run(run, size_px, color, bold=False, tracking_em=0):
    run.font.name = "Poppins"
    run.font.size = Pt(size_px / 2)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = rgb(color)
    r_pr = run._r.get_or_add_rPr()
    r_pr.set("spc", str(int(round(tracking_em * (size_px / 2) * 100))))
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = r_pr.find(qn(tag))
        if el is None:
            el = etree.SubElement(r_pr, qn(tag))
        el.set("typeface", "Poppins")


def add_round(slide, x, y, w, h, fill, radius):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    short = min(w, h) or 1
    shape.adjustments[0] = min(0.5, radius / short)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, size, color, bold=False, tracking_em=0, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    set_run(run, size, color, bold, tracking_em)
    return frame


def apply_list_bullet(paragraph, gap_px):
    """One real list paragraph. Enter in Google Slides continues the bullet."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(int(round(28 * EMU_PER_PX))))
    p_pr.set("indent", str(int(round(-22 * EMU_PER_PX))))
    p_pr.set("lvl", "0")
    for child in list(p_pr):
        p_pr.remove(child)

    def append(parent, tag, **attrs):
        element = etree.SubElement(parent, qn(tag))
        for key, value in attrs.items():
            element.set(key, value)
        return element

    append(append(p_pr, "a:lnSpc"), "a:spcPct", val="120000")
    append(append(p_pr, "a:spcAft"), "a:spcPts", val=str(int(round(gap_px / 2 * 100))))
    append(append(p_pr, "a:buClr"), "a:srgbClr", val="395AFA")
    append(p_pr, "a:buSzPct", val="80000")
    append(p_pr, "a:buFont", typeface="Arial")
    append(p_pr, "a:buChar", char="•")
    default = append(p_pr, "a:defRPr", sz="1400", dirty="0")
    append(append(default, "a:solidFill"), "a:srgbClr", val="1B1A18")
    append(default, "a:latin", typeface="Poppins")
    append(default, "a:ea", typeface="Poppins")
    append(default, "a:cs", typeface="Poppins")


def add_bullet_list(slide, x, y, w, h, items, ink, gap_px):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, margin, 0)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        if isinstance(item, tuple):
            label, text = item
            rest = PLACE if is_prompt(text) else ink
            lead = paragraph.add_run()
            lead.text = f"{label}: "
            set_run(lead, 28, ink, bold=True)
            tail = paragraph.add_run()
            tail.text = text
            set_run(tail, 28, rest, bold=False)
        else:
            run = paragraph.add_run()
            run.text = item
            set_run(run, 28, PLACE if is_prompt(item) else ink, bold=False)
        apply_list_bullet(paragraph, gap_px)
    return frame


def build_pptx(spec, out_path):
    body = font(poppins(400), 28, 400)
    strong = font(poppins(600), 28, 600)
    assert_fits(spec, body, strong)
    presentation = Presentation()
    presentation.slide_width = emu(W)
    presentation.slide_height = emu(H)
    presentation.core_properties.title = f"Growth Play: {spec['title']}"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(OAT)

    add_text(slide, PAD_X, 62, 1100, 28, "GROWTH PLAY", 20, MUTED, bold=True, tracking_em=0.12)
    add_text(slide, PAD_X, 96, 1280, 84, spec["title"], 80, BLACK, bold=True, tracking_em=-0.04)

    live_label = font(poppins(600), 18, 600)
    live_date = font(poppins(500), 26, 500)
    pill_w = 26 + tracked_width("GO-LIVE", live_label, 0.12) + 14 + live_date.getlength(spec["go_live"]) + 26
    pill_h = 59
    pill_x = W - PAD_X - pill_w
    pill_y = HEADER_BOTTOM - 4 - pill_h
    add_round(slide, pill_x, pill_y, pill_w, pill_h, SLAB, pill_h / 2)
    add_text(slide, pill_x + 22, pill_y + 12, 140, 36, "GO-LIVE", 18, MUTED, bold=True, tracking_em=0.12, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, pill_x + 22 + 130, pill_y + 10, pill_w - 170, 40, spec["go_live"], 26, INK, anchor=MSO_ANCHOR.MIDDLE)

    for panel in PANELS:
        x, y, h = panel["x"], panel["y"], panel["h"]
        add_round(slide, x, y, COL_W, h, panel["fill"], 24)
        slide.shapes.add_picture(str(ASSETS / "icons" / panel["icon"]), emu(x + 40), emu(y + 32), emu(80), emu(80))
        add_text(slide, x + 140, y + 48, COL_W - 180, 48, panel["heading"], 40, panel["ink"], bold=True, tracking_em=-0.025, anchor=MSO_ANCHOR.MIDDLE)
        add_bullet_list(
            slide,
            x + 40,
            y + 134,
            COL_W - 80,
            h - 154,
            content_of(spec, panel),
            panel["ink"],
            panel["gap"],
        )

    logo = Image.open(ASSETS / "logo" / "Clay_Logo_3D_Blk.png")
    logo_h = 40
    logo_w = round(logo.width * logo_h / logo.height)
    slide.shapes.add_picture(
        str(ASSETS / "logo" / "Clay_Logo_3D_Blk.png"),
        emu(PAD_X),
        emu(1000 + (120 - 40) / 2),
        emu(logo_w),
        emu(logo_h),
    )
    presentation.save(out_path)


def output_paths(spec, out_dir):
    safe = spec["title"].replace("/", "-")
    stem = f"Growth Play - {safe}"
    return out_dir / f"{stem}.pptx", out_dir / f"{stem}.png"


def main():
    parser = argparse.ArgumentParser(description="Generate a Growth Play slide.")
    parser.add_argument("spec", help="Path to the slide JSON")
    parser.add_argument("--out", default=str(Path.home() / "Downloads"), help="Output directory")
    parser.add_argument("--open", action="store_true", help="Open the PNG preview")
    args = parser.parse_args()
    spec = load_spec(args.spec)
    out_dir = Path(args.out).expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    pptx_path, png_path = output_paths(spec, out_dir)
    roobert = ensure_roobert()
    if roobert:
        sans, mono = roobert
        preview = "Roobert"
    else:
        sans, mono = poppins(400), poppins(600)
        preview = "Poppins"
    build_pptx(spec, pptx_path)
    build_png(spec, png_path, sans, mono)
    print(pptx_path)
    print(png_path)
    print(f"preview font: {preview}")
    if args.open:
        subprocess.run(["open", str(png_path)], check=False)


if __name__ == "__main__":
    try:
        main()
    except BrokenPipeError:
        sys.exit(0)
